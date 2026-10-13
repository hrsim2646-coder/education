# -*- coding: utf-8 -*-
"""samples.json -> (A) 웹앱 HTML (B) 엑셀 (C) 워드  /  실행: python -I build_outputs.py"""
import json, os, math
HERE = os.path.dirname(os.path.abspath(__file__))
DRAFT = os.path.join(HERE, '..', 'DRAFT')
DATA = json.load(open(os.path.join(DRAFT, 'samples.json'), encoding='utf-8'))
META = {
    'basis': '관세법 2026.10.2 시행 · 시행령 2026.7.1 시행 · 시행규칙 2026.7.31 시행',
    'sources': ['출처: 국가법령정보센터 PDF (관세법 제21490호 / 시행령 제36140호 / 시행규칙 제42호)',
                '※ 학습·정리용 초안입니다. 연계(위임/인용) 분류는 본문 인용 패턴으로 자동 추출한 결과이므로 최종 확인은 원문으로 하세요.'],
}
KL = {'law': '법', 'decree': '령', 'rule': '칙'}
KF = {'law': '관세법', 'decree': '시행령', 'rule': '시행규칙'}

def ptext(c, limit=None, sep='\n'):
    out = []
    for p in c['paras']:
        out.append(('    ' * max(p['l'] - 1, 0)) + p['t'])
    t = sep.join(out)
    if limit and len(t) > limit:
        t = t[:limit].rstrip() + ' … (이하 생략, 전문은 조문DB)'
    return t

def label(c):
    return f"{KF[c['kind']]} 제{c['no']}조({c['title']})"

# ───────────────────────── A. 웹앱 ─────────────────────────
def build_web():
    tpl = open(os.path.join(HERE, 'web_template.html'), encoding='utf-8').read()
    html = tpl.replace('/*DATA*/null', json.dumps(DATA, ensure_ascii=False)).replace('/*META*/null', json.dumps(META, ensure_ascii=False))
    p = os.path.join(DRAFT, '관세법_연계뷰_웹앱.html')
    open(p, 'w', encoding='utf-8').write(html)
    return p

# ───────────────────────── B. 엑셀 ─────────────────────────
def build_xlsx():
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    F = {'law': 'DCE8FB', 'decree': 'DFF2E7', 'rule': 'FCE9D6', 'head': '1F2937', 'gos': 'ECEEF1'}
    thin = Side(style='thin', color='C9D0DA')
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    wrap = Alignment(wrap_text=True, vertical='top')
    font = Font(name='맑은 고딕', size=10)
    bold = Font(name='맑은 고딕', size=10, bold=True)

    def head(ws, row, titles, widths):
        for i, (t, w) in enumerate(zip(titles, widths), 1):
            c = ws.cell(row=row, column=i, value=t)
            c.font = Font(name='맑은 고딕', size=10, bold=True, color='FFFFFF')
            c.fill = PatternFill('solid', fgColor=F['head'])
            c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            c.border = box
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.row_dimensions[row].height = 26

    def height(texts_widths):
        lines = 1
        for t, w in texts_widths:
            n = 0
            for seg in str(t).split('\n'):
                n += max(1, math.ceil(len(seg) * 1.9 / max(w, 1)))
            lines = max(lines, n)
        return min(409, 14.5 * lines + 4)

    # 1) 안내
    ws = wb.active; ws.title = '안내'
    ws.column_dimensions['A'].width = 22; ws.column_dimensions['B'].width = 90
    rows = [
        ('제목', '관세법 법·령·칙 연계표 (초안)'),
        ('기준', META['basis']),
        ('시트 구성', '① 3단연계표 — 법 조문을 기준으로 시행령·시행규칙을 한 행에 이어서 보기\n② 조문DB — 조문 1개 = 1행(원문 전체). 후속 문제은행·논술·웹앱의 원천 데이터'),
        ('연계 구분', '위임: 하위 조문이 상위 조문을 근거로 직접 인용(항 첫머리·제목 유사)\n인용: 호·목 안에서 단순 언급(위임 근거는 아님)'),
        ('색상', '파랑=법 / 초록=령 / 주황=칙 / 회색=고시'),
        ('주의', '연계는 본문 인용 패턴으로 자동 추출한 초안입니다. 고시는 원문 PDF 미적재 상태이며 “고시 포함” 여부만 표시합니다.'),
        ('후속 활용', '조문DB의 ID(L-10, D-2, R-4의3)를 기준으로 객관식·논술·암기카드를 자동 생성할 수 있습니다.'),
    ]
    for i, (a, b) in enumerate(rows, 1):
        ws.cell(i, 1, a).font = bold; ws.cell(i, 1).fill = PatternFill('solid', fgColor='EEF1F6')
        c = ws.cell(i, 2, b); c.font = font; c.alignment = wrap
        ws.cell(i, 1).alignment = Alignment(vertical='top')
        ws.row_dimensions[i].height = 18 * (b.count('\n') + 1) + 6

    # 2) 3단연계표
    ws = wb.create_sheet('3단연계표')
    titles = ['법 조문', '시행령 조문', '령 구분', '시행규칙 조문', '칙 구분', '고시', '법 본문', '시행령 본문', '시행규칙 본문']
    widths = [22, 24, 8, 24, 8, 8, 52, 52, 52]
    head(ws, 1, titles, widths)
    r = 2
    for c in DATA:
        L = c['law']; start = r
        main = [d for d in c['decrees'] if d['link'] == '위임']
        ref = [d for d in c['decrees'] if d['link'] != '위임']
        lines = []  # (decree, rule)
        for d in main + ref:
            if d['rules']:
                for ru in d['rules']: lines.append((d, ru))
            else:
                lines.append((d, None))
        for ru in c['rules_direct']: lines.append((None, ru))
        if not lines: lines.append((None, None))
        gos = '●' if any('고시' in x['deleg'] for x in [L] + [y for y, _ in lines if y] + [z for _, z in lines if z]) else ''
        first = True
        for d, ru in lines:
            vals = [
                f"제{L['no']}조\n{L['title']}" if first else None,
                label(d).replace('시행령 ', '') if d else '(직접 위임)' if ru else '—',
                d['link'] if d else '',
                label(ru).replace('시행규칙 ', '') if ru else '',
                ru['link'] if ru else '',
                gos if first else '',
                ptext(L, 700) if first else None,
                ptext(d, 600) if d else '',
                ptext(ru, 500) if ru else '',
            ]
            for i, v in enumerate(vals, 1):
                cell = ws.cell(r, i, v); cell.font = font; cell.alignment = wrap; cell.border = box
            ws.cell(r, 1).fill = PatternFill('solid', fgColor=F['law']); ws.cell(r, 1).font = bold
            ws.cell(r, 7).fill = PatternFill('solid', fgColor=F['law'])
            for col in (2, 3, 8): ws.cell(r, col).fill = PatternFill('solid', fgColor=F['decree'] if (d and d['link'] == '위임') else 'F4F9F6')
            for col in (4, 5, 9): ws.cell(r, col).fill = PatternFill('solid', fgColor=F['rule'] if (ru and ru['link'] == '위임') else 'FDF6EE')
            ws.cell(r, 6).fill = PatternFill('solid', fgColor=F['gos']); ws.cell(r, 6).alignment = Alignment(horizontal='center', vertical='top')
            for col in (3, 5): ws.cell(r, col).alignment = Alignment(horizontal='center', vertical='top')
            ws.row_dimensions[r].height = height([(vals[6] or '', 52), (vals[7], 52), (vals[8], 52)]) if first else height([(vals[7], 52), (vals[8], 52)])
            first = False
            r += 1
        end = r - 1
        if end > start:
            for col in (1, 6, 7):
                ws.merge_cells(start_row=start, start_column=col, end_row=end, end_column=col)
            # 법 본문은 병합 셀이라 높이를 행 합계로 맞춤
        # 법 본문이 긴 경우 병합영역 높이 보정
        need = height([(ptext(L, 700), 52)])
        have = sum(ws.row_dimensions[x].height or 15 for x in range(start, end + 1))
        if have < need:
            add = (need - have) / (end - start + 1)
            for x in range(start, end + 1): ws.row_dimensions[x].height = min(409, (ws.row_dimensions[x].height or 15) + add)
    ws.freeze_panes = 'C2'
    ws.auto_filter.ref = f"A1:{get_column_letter(len(titles))}{r-1}"
    ws.sheet_view.zoomScale = 90

    # 3) 조문DB
    ws = wb.create_sheet('조문DB')
    titles = ['ID', '구분', '조번호', '제목', '편·장', '절·관', '상위 ID', '연계', '위임 대상', '본문(전문)']
    widths = [10, 8, 9, 30, 20, 20, 10, 8, 12, 90]
    head(ws, 1, titles, widths)
    seen = {}
    def add(c, parent=None, link=''):
        cid = f"{ {'law':'L','decree':'D','rule':'R'}[c['kind']] }-{c['no']}"
        if cid in seen: return
        seen[cid] = c
        deleg = ','.join({'령': '령', '칙': '칙', '고시': '고시'}[x] for x in c['deleg'])
        body = ptext(c)
        rown = ws.max_row + 1
        vals = [cid, KF[c['kind']], f"제{c['no']}조", c['title'], c['chapter'], c['section'], parent or '', link, deleg, body]
        for i, v in enumerate(vals, 1):
            cell = ws.cell(rown, i, v); cell.font = font; cell.alignment = wrap; cell.border = box
        ws.cell(rown, 1).fill = PatternFill('solid', fgColor=F[c['kind']]); ws.cell(rown, 2).fill = PatternFill('solid', fgColor=F[c['kind']])
        ws.row_dimensions[rown].height = height([(body, 90), (c['title'], 30)])
    for c in DATA:
        L = c['law']; lid = f"L-{L['no']}"
        add(L)
        for d in c['decrees']:
            add(d, lid, d['link'])
            for ru in d['rules']: add(ru, f"D-{d['no']}", ru['link'])
        for ru in c['rules_direct']: add(ru, lid, ru['link'])
    ws.freeze_panes = 'E2'
    ws.auto_filter.ref = f"A1:{get_column_letter(len(titles))}{ws.max_row}"
    p = os.path.join(DRAFT, '관세법_연계표_엑셀.xlsx')
    wb.save(p)
    return p

# ───────────────────────── C. 워드 ─────────────────────────
def build_docx():
    from docx import Document
    from docx.shared import Pt, Cm, RGBColor
    from docx.enum.section import WD_ORIENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
    for m in ('left_margin', 'right_margin'): setattr(sec, m, Cm(1.4))
    sec.top_margin = Cm(1.3); sec.bottom_margin = Cm(1.3)

    def setfont(run, size=9, bold=False, color=None):
        run.font.name = '맑은 고딕'; run.font.size = Pt(size); run.font.bold = bold
        run._element.rPr.rFonts.set(qn('w:eastAsia'), '맑은 고딕')
        if color: run.font.color.rgb = RGBColor.from_string(color)

    st = doc.styles['Normal']; st.font.name = '맑은 고딕'; st.font.size = Pt(9)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), '맑은 고딕')

    def shade(cell, hex_):
        tcPr = cell._element.get_or_add_tcPr()
        shd = OxmlElement('w:shd'); shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), hex_)
        tcPr.append(shd)

    def para(cell_or_doc, text='', size=9, bold=False, color=None, first=False, indent=0, after=2):
        p = (cell_or_doc.paragraphs[0] if cell_or_doc.paragraphs else cell_or_doc.add_paragraph()) if first else cell_or_doc.add_paragraph()
        p.paragraph_format.space_after = Pt(after); p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.line_spacing = 1.15
        if indent: p.paragraph_format.left_indent = Cm(0.45 * indent)
        if text:
            setfont(p.add_run(text), size, bold, color)
        return p

    def fill_card(cell, c, limit=900):
        col = {'law': '1F5FBF', 'decree': '1B7F4B', 'rule': 'C2570C'}[c['kind']]
        link = f" [{c['link']}]" if c.get('link') else ''
        para(cell, f"{KF[c['kind']]} 제{c['no']}조({c['title']}){link}", 9.5, True, col, first=True, after=3)
        used = 0
        for p in c['paras']:
            t = p['t']
            if used + len(t) > limit:
                para(cell, t[:max(0, limit - used)].rstrip() + ' … (이하 생략)', 8.5, indent=max(p['l'] - 1, 0)); break
            para(cell, t, 8.5, indent=max(p['l'] - 1, 0)); used += len(t)

    # 표지
    para(doc, '관세법 법·령·칙 3단 연계 대비표', 20, True, '1F2937', first=True, after=2)
    para(doc, META['basis'], 10, False, '6B7685', after=2)
    para(doc, '조문 하나를 기준으로 [법] → [시행령] → [시행규칙]을 한 장에 이어서 봅니다. 파랑=법, 초록=령, 주황=칙, 회색=고시.', 9.5, after=8)

    # 요약표
    t = doc.add_table(rows=1, cols=4); t.style = 'Table Grid'
    for i, h in enumerate(['법 조문', '시행령 (위임)', '시행규칙', '고시']):
        cell = t.rows[0].cells[i]; shade(cell, '1F2937'); para(cell, h, 9.5, True, 'FFFFFF', first=True)
    for c in DATA:
        L = c['law']; rc = t.add_row().cells
        main = [d for d in c['decrees'] if d['link'] == '위임']
        rules = [r for d in main for r in d['rules']] + c['rules_direct']
        gos = any('고시' in x['deleg'] for x in [L] + main + rules)
        shade(rc[0], 'DCE8FB'); shade(rc[1], 'DFF2E7'); shade(rc[2], 'FCE9D6'); shade(rc[3], 'ECEEF1')
        para(rc[0], f"제{L['no']}조 {L['title']}", 9, True, first=True)
        para(rc[1], '\n'.join(f"제{d['no']}조 {d['title']}" for d in main) or '—', 9, first=True)
        para(rc[2], '\n'.join(f"제{r['no']}조 {r['title']}" for r in rules) or '—', 9, first=True)
        para(rc[3], '● 포함' if gos else '—', 9, first=True)
    for row in t.rows:
        for w, cell in zip((Cm(6.2), Cm(8.2), Cm(8.2), Cm(2.4)), row.cells): cell.width = w

    # 조문별 3단 표
    for c in DATA:
        doc.add_page_break()
        L = c['law']
        para(doc, f"법 제{L['no']}조 {L['title']}", 15, True, '1F5FBF', first=True, after=0)
        loc = ' › '.join(x for x in (L['chapter'], L['section']) if x)
        para(doc, loc, 9, False, '6B7685', after=6)
        main = [d for d in c['decrees'] if d['link'] == '위임']
        ref = [d for d in c['decrees'] if d['link'] != '위임']
        nrows = max(1, len(main) + (1 if c['rules_direct'] else 0))
        tb = doc.add_table(rows=1 + nrows, cols=3); tb.style = 'Table Grid'; tb.autofit = False
        for i, (h, colr) in enumerate((('법 (관세법)', '1F5FBF'), ('시행령', '1B7F4B'), ('시행규칙', 'C2570C'))):
            cell = tb.rows[0].cells[i]; shade(cell, colr); para(cell, h, 10, True, 'FFFFFF', first=True)
        lc = tb.cell(1, 0).merge(tb.cell(nrows, 0)); shade(lc, 'EEF4FD'); fill_card(lc, L, 1800)
        ri = 1
        for d in main:
            cd, cr = tb.cell(ri, 1), tb.cell(ri, 2)
            shade(cd, 'F1FAF5'); fill_card(cd, d, 1100)
            shade(cr, 'FFF7EE')
            if d['rules']:
                for k, rr in enumerate(d['rules']):
                    if k: para(cr, '', 4)
                    if k == 0:
                        fill_card(cr, rr, 800)
                    else:
                        p = para(cr, f"{KF['rule']} 제{rr['no']}조({rr['title']}) [{rr['link']}]", 9.5, True, 'C2570C')
                        for pp in rr['paras']: para(cr, pp['t'][:500], 8.5, indent=max(pp['l'] - 1, 0))
            else:
                para(cr, '연결된 시행규칙 조문 없음', 8.5, False, '9AA3B0', first=True)
            ri += 1
        if c['rules_direct']:
            cd, cr = tb.cell(ri, 1), tb.cell(ri, 2)
            shade(cd, 'F4F5F7'); para(cd, '법 → 시행규칙 직접 위임', 9, True, '6B7685', first=True)
            para(cd, '시행령을 거치지 않고 법에서 바로 부령으로 위임', 8.5, False, '6B7685')
            shade(cr, 'FFF7EE')
            for k, rr in enumerate(c['rules_direct']):
                if k: para(cr, '', 4)
                if k == 0: fill_card(cr, rr, 800)
                else:
                    para(cr, f"{KF['rule']} 제{rr['no']}조({rr['title']}) [{rr['link']}]", 9.5, True, 'C2570C')
                    for pp in rr['paras']: para(cr, pp['t'][:500], 8.5, indent=max(pp['l'] - 1, 0))
        if not main and not c['rules_direct']:
            para(tb.cell(1, 1), '이 조문을 위임 근거로 삼는 시행령 조문 없음', 8.5, False, '9AA3B0', first=True)
            para(tb.cell(1, 2), '—', 8.5, False, '9AA3B0', first=True)
        for row in tb.rows:
            for w, cell in zip((Cm(8.4), Cm(9.0), Cm(8.6)), row.cells): cell.width = w
        if ref:
            para(doc, '', 4)
            para(doc, '▸ 단순 인용 조문(위임 근거는 아니나 이 조문을 언급): ' + ', '.join(f"령 제{d['no']}조({d['title']})" for d in ref), 8.5, False, '6B7685')
        if any('고시' in x['deleg'] for x in [L] + c['decrees']):
            para(doc, '▸ 고시: 본문에 “관세청장이 정하여 고시” 위임이 포함되어 있음 (고시 원문 미적재)', 8.5, False, '6B7685')

    p = os.path.join(DRAFT, '관세법_3단대비표_워드.docx')
    doc.save(p)
    return p

if __name__ == '__main__':
    for fn in (build_web, build_xlsx, build_docx):
        print(fn.__name__, '->', os.path.basename(fn()))
